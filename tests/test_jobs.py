import asyncio
from dataclasses import replace
import io
import json
from pathlib import Path
import tempfile
import time
import unittest
from types import SimpleNamespace
from unittest.mock import AsyncMock, patch
import uuid

from fastapi.testclient import TestClient
from openpyxl import load_workbook

import app as web
from alignment import AlignmentConfig
from jobs import JobManager, RETENTION_SECONDS, read_state, write_state
from workbook_limits import MAX_UPLOAD_BYTES, inspect_workbook
from workbook_fixture import PHOTO_OPTIONS, make_workbook


class JobApiTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.data_patch = patch.object(web, "DATA_DIR", self.root / "data")
        self.db_patch = patch.object(web, "USAGE_DB", self.root / "data" / "usage.db")
        self.data_patch.start()
        self.db_patch.start()
        self.client = TestClient(web.app)
        self.client.__enter__()
        self.client.cookies.set(web.AUTH_COOKIE_NAME, web.auth_cookie_value())

    def tearDown(self):
        self.client.__exit__(None, None, None)
        self.db_patch.stop()
        self.data_patch.stop()
        self.temp.cleanup()

    def upload(self, path, **options):
        with path.open("rb") as file:
            return self.client.post("/process", files={"file": (path.name, file)}, data={**PHOTO_OPTIONS, **options})

    def wait(self, job_id, timeout=90):
        deadline = time.monotonic() + timeout
        while time.monotonic() < deadline:
            status = self.client.get(f"/jobs/{job_id}")
            self.assertEqual(status.status_code, 200, status.text)
            result = status.json()
            if result["status"] in {"succeeded", "failed"}:
                return result
            time.sleep(0.05)
        self.fail("Job did not complete within the test deadline")

    def test_10000_styled_rows_background_download_and_responsiveness(self):
        source = self.root / "ten-thousand.xlsx"
        make_workbook(source)
        started = time.monotonic()
        response = self.upload(source)
        self.assertEqual(response.status_code, 202, response.text)
        job = response.json()
        self.assertEqual(job["workflow"], "large")
        self.assertEqual(job["rows"], 10_000)
        self.assertEqual(self.client.get(job["download_url"]).status_code, 409)
        # Exercise the same web server while the real CPU-heavy worker is running.
        latencies = []
        saw_processing = False
        deadline = time.monotonic() + 90
        while time.monotonic() < deadline:
            before = time.monotonic()
            health = self.client.get("/healthz")
            current = self.client.get(job["status_url"])
            latencies.append(time.monotonic() - before)
            self.assertEqual(health.status_code, 200)
            job = current.json()
            saw_processing |= job["status"] == "processing"
            if job["status"] in {"succeeded", "failed"}:
                break
            time.sleep(0.1)
        self.assertTrue(saw_processing)
        self.assertEqual(job["status"], "succeeded", job)
        self.assertLess(max(latencies), 2, "Status and health requests must remain responsive")
        downloaded = self.client.get(job["download_url"])
        self.assertEqual(downloaded.status_code, 200)
        self.assertIn("ten-thousand_shifted.xlsx", downloaded.headers["content-disposition"])
        # A fresh client can reconnect to the same accepted job without uploading.
        fresh = TestClient(web.app)
        fresh.cookies.set(web.AUTH_COOKIE_NAME, web.auth_cookie_value())
        self.assertEqual(fresh.get(job["status_url"]).json()["status"], "succeeded")
        fresh.close()
        workbook = load_workbook(io.BytesIO(downloaded.content))
        original = workbook["sheet1"]
        aligned = workbook["Aligned results"]
        rows = list(aligned.iter_rows(min_row=3, values_only=True))
        self.assertEqual([row[3] for row in rows if row[3] is not None], list(range(10_000)))
        self.assertEqual([row[10] for row in rows if row[10] is not None], list(range(4312)) + list(range(4313, 10_001)))
        self.assertEqual([row[4] for row in rows if row[3] is not None], [f"=D{i+3}*2" for i in range(10_000)])
        self.assertEqual(aligned.max_row, 10_003)  # Two headers plus inserted alignment gaps.
        self.assertEqual(aligned["G3"]._style, original["G3"]._style)
        self.assertEqual(aligned.column_dimensions["G"].width, 24)
        self.assertEqual(original.max_row, 10_002)
        self.assertAlmostEqual(aligned["H3"].value, 0.1)
        workbook.close()
        print(f"\n10,000 styled rows: {time.monotonic()-started:.2f}s; worst health+status response {max(latencies):.3f}s")

    def test_small_file_direct_download(self):
        path = self.root / "small.xlsx"
        make_workbook(path, 8)
        response = self.upload(path)
        self.assertEqual(response.status_code, 200, response.text[:200])
        self.assertTrue(response.content.startswith(b"PK"))

    def test_10001_rows_rejected_without_processing_or_truncation(self):
        path = self.root / "too-many.xlsx"
        make_workbook(path, 10_001, styled=False)
        response = self.upload(path)
        self.assertEqual(response.status_code, 400)
        self.assertIn("10,000", response.json()["detail"])
        self.assertEqual(len(web.app.state.jobs.active), 0)
        self.assertFalse(list((self.root / "data/jobs").glob("*/input.xlsx")))

    def test_explicit_large_mode_and_duplicate_upload_id(self):
        path = self.root / "small.xlsx"
        make_workbook(path, 4)
        job_id = uuid.uuid4().hex
        def send():
            with path.open("rb") as file:
                return self.client.post("/process", headers={"X-Upload-ID": job_id},
                    files={"file": (path.name, file)}, data={**PHOTO_OPTIONS, "workflow": "large"})
        first, second = send(), send()
        self.assertEqual(first.status_code, 202)
        self.assertEqual(second.json()["id"], first.json()["id"])
        self.assertEqual(self.wait(job_id)["status"], "succeeded")
        self.assertEqual(self.client.get("/usage").json()["runs"], 1)
        self.assertEqual(len(list((self.root / "data/jobs").glob("*/status.json"))), 1)

    def test_failure_releases_queue_and_next_job_succeeds(self):
        path = self.root / "empty-column.xlsx"
        make_workbook(path, 3)
        response = self.upload(path, workflow="large", left_input_col="A")
        self.assertEqual(response.status_code, 202)
        failed = self.wait(response.json()["id"])
        self.assertEqual(failed["status"], "failed")
        self.assertIn("numeric", failed["error"])
        success = self.upload(path, workflow="large")
        self.assertEqual(self.wait(success.json()["id"])["status"], "succeeded")

    def test_auth_and_missing_job(self):
        self.assertEqual(self.client.get("/jobs/" + "a" * 32).status_code, 404)
        self.client.cookies.clear()
        self.assertEqual(self.client.get("/jobs/" + "a" * 32).status_code, 401)
        self.assertEqual(self.client.get("/jobs/" + "a" * 32 + "/download").status_code, 401)
        self.assertEqual(self.client.get("/healthz").status_code, 200)

    def test_queue_busy_and_upload_byte_limit(self):
        manager = web.app.state.jobs
        for _ in range(3):
            manager.reserve(uuid.uuid4().hex, "test.xlsx")
        path = self.root / "small.xlsx"
        make_workbook(path, 2)
        self.assertEqual(self.upload(path).status_code, 429)
        response = self.client.post("/process", content=b"x", headers={"Content-Length": str(MAX_UPLOAD_BYTES + 100_000)})
        self.assertEqual(response.status_code, 413)

    def test_worker_timeout_is_reported_and_worker_recovers(self):
        manager = web.app.state.jobs
        manager.timeout = 0.001
        path = self.root / "small.xlsx"
        make_workbook(path, 2)
        response = self.upload(path, workflow="large")
        failed = self.wait(response.json()["id"])
        self.assertEqual(failed["status"], "failed")
        self.assertIn("exceeded", failed["error"])
        manager.timeout = 600
        response = self.upload(path, workflow="large")
        self.assertEqual(self.wait(response.json()["id"])["status"], "succeeded")

    def test_chunked_upload_cannot_bypass_size_limit(self):
        def body():
            yield b'--test-boundary\r\nContent-Disposition: form-data; name="file"; filename="big.xlsx"\r\n\r\n'
            yield b"x" * (70 * 1024)
            yield b"\r\n--test-boundary--\r\n"
        with patch.object(web, "MAX_UPLOAD_BYTES", 1024):
            response = self.client.post("/process", content=body(), headers={"Content-Type": "multipart/form-data; boundary=test-boundary"})
        self.assertEqual(response.status_code, 413, response.text)

    def test_lost_worker_does_not_leave_job_processing_forever(self):
        manager = web.app.state.jobs
        path = self.root / "small.xlsx"
        make_workbook(path, 2)
        # The worker exits before it can write a result (e.g. killed by the host).
        dead = SimpleNamespace(returncode=-9, wait=AsyncMock(return_value=-9))
        with patch("jobs.asyncio.create_subprocess_exec", new=AsyncMock(return_value=dead)):
            response = self.upload(path, workflow="large")
            failed = self.wait(response.json()["id"])
            self.assertEqual(failed["status"], "failed")
            self.assertIn("stopped unexpectedly", failed["error"])
        response = self.upload(path, workflow="large")
        self.assertEqual(self.wait(response.json()["id"])["status"], "succeeded")


class LimitTests(unittest.TestCase):
    def test_trailing_formatting_does_not_count_as_data(self):
        with tempfile.TemporaryDirectory() as temporary:
            path = Path(temporary) / "styled-tail.xlsx"
            make_workbook(path, 2)
            workbook = load_workbook(path)
            workbook.active["A1048576"]._style = workbook.active["A2"]._style
            workbook.save(path)
            self.assertEqual(inspect_workbook(path, AlignmentConfig(**PHOTO_OPTIONS))["rows"], 2)

    def test_second_input_sheet_is_also_limited(self):
        with tempfile.TemporaryDirectory() as temporary:
            path = Path(temporary) / "two-sheets.xlsx"
            make_workbook(path, 2)
            workbook = load_workbook(path)
            workbook.create_sheet("Other")["I10003"] = 1
            workbook.save(path)
            with self.assertRaisesRegex(ValueError, "10,000"):
                inspect_workbook(path, AlignmentConfig(**PHOTO_OPTIONS, right_sheet_name="Other"))

    def test_invalid_files_and_invalid_settings(self):
        with tempfile.TemporaryDirectory() as temporary:
            path = Path(temporary) / "invalid.xlsx"
            path.write_text("not a workbook")
            config = AlignmentConfig(**PHOTO_OPTIONS)
            with self.assertRaisesRegex(ValueError, "could not be read"):
                inspect_workbook(path, config)
            with self.assertRaisesRegex(ValueError, "finite"):
                inspect_workbook(path, replace(config, threshold=float("nan")))


class LifecycleTests(unittest.IsolatedAsyncioTestCase):
    async def test_second_web_worker_is_rejected(self):
        with tempfile.TemporaryDirectory() as temporary:
            first = JobManager(temporary, lambda: 1)
            second = JobManager(temporary, lambda: 1)
            await first.start()
            try:
                with self.assertRaisesRegex(RuntimeError, "one Uvicorn worker"):
                    await second.start()
            finally:
                await first.stop()

    async def test_restart_and_expiration_cleanup(self):
        with tempfile.TemporaryDirectory() as temporary:
            manager = JobManager(temporary, lambda: 1)
            job_id = uuid.uuid4().hex
            manager.reserve(job_id, "test.xlsx")
            (manager.directory(job_id) / "input.xlsx").write_bytes(b"unfinished")
            restarted = JobManager(temporary, lambda: 1)
            await restarted.start()
            try:
                self.assertEqual(restarted.status(job_id)["status"], "failed")
                self.assertIn("restarted", restarted.status(job_id)["error"])
                self.assertFalse((restarted.directory(job_id) / "input.xlsx").exists())
                state = restarted.status(job_id)
                state["updated_at"] = time.time() - RETENTION_SECONDS - 1
                write_state(restarted.directory(job_id), state)
                await restarted.cleanup()
                self.assertFalse(restarted.directory(job_id).exists())
            finally:
                await restarted.stop()


if __name__ == "__main__":
    unittest.main()
