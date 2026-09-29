"""Generate a realistic styled workbook; no customer data is checked in."""
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

PHOTO_OPTIONS = {
    "input_sheet_name": "sheet1", "output_sheet_name": "Aligned results",
    "start_row": 3, "header_first_row": 1, "header_last_row": 2,
    "left_input_col": "G", "left_block_start_col": "A", "left_block_end_col": "G",
    "left_output_start_col": "A", "right_input_col": "I", "right_block_start_col": "I",
    "right_block_end_col": "O", "right_output_start_col": "I", "threshold": 0.5,
    "diff_output_col": "H",
}


def make_workbook(path, rows=10_000, styled=True):
    wb = Workbook()
    ws = wb.active
    ws.title = "sheet1"
    ws.append(["Current", None, None, None, None, None, None, None, "Previous"])
    ws.append(["Seam", "Type", "WT", "Joint #", "Formula", "Distance", "Target Joint Length [ft]",
               "Difference", "Joint Length [ft]", "Distance", "Joint #", "WT", "Type", "Seam", "Note"])
    template = ws["A2"]
    template.font = Font(name="Calibri", size=11, bold=True, color="123456")
    template.fill = PatternFill("solid", fgColor="E2EFDA")
    template.border = Border(bottom=Side(style="thin", color="AAAAAA"))
    template.alignment = Alignment(horizontal="right")
    template.number_format = '0.000'
    for index in range(rows):
        right_index = index if index < 4312 else index + 1
        ws.append(["LS", "Type", 0.219, index, f"=D{index + 3}*2", index * 2.5,
                   10 + index * 2.5, None, 10.1 + right_index * 2.5,
                   right_index * 2.5, right_index, 0.219, "Type", "LS", ""])
        if styled:
            for column in range(1, 16):
                ws.cell(index + 3, column)._style = template._style
    ws.column_dimensions["G"].width = 24
    wb.save(path)
    wb.close()
